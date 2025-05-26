from flask import Flask, request, jsonify
from flask_cors import CORS
import google.generativeai as genai
import os
from dotnev import load_dotenv
import fitz  #For PDF
import docx  #For DOCX
import openpyxl  #For Excel

load_dotenv()

app = Flask(__name__)
CORS(app)

genai.configure(api_key=os.getenv("GEMINI_API_KEY"))
model = genai.GenerativeModel("gemini-2.0-flash")

def extract_text(file_storage):
    filename =file_storage.filename
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        return file_storage.read().decode("utf-8")

    elif ext == ".pdf":
        pdf_text = ""
        doc=fitz.open(stream=file_storage.read(), filetype="pdf")
        for page in doc:
            pdf_text += page.get_text()
        return pdf_text

    elif ext == ".docx":
        doc= docx.Document(file_storage)
        return "\n".join([p.text for p in doc.paragraphs])

    elif ext == ".xlsx":
        workbook=openpyxl.load_workbook(file_storage, data_only=True)
        text = ""
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                line= " ".join([str(cell) for cell in row if cell is not None])
                text += line + "\n"
        return text

    else:
        return None

@app.route("/parse", methods=["POST"])
def parse_resume():
    if "file" not in request.files:
        return jsonify({"result": "No file uploaded"}),400

    file=request.files["file"]
    text=extract_text(file)

    if not text:
        return jsonify({"result": "Unsupported file format or empty file"}),400

    prompt = f"""
    Extract the following details from the resume:
    - Name
    - Email
    - Phone
    - Skills
    - Experience
    - Education

    Resume Content:
    {text}
    """

    try:
        response = model.generate_content(prompt)
        return jsonify({"result": response.text.strip()})
    except Exception as e:
        return jsonify({"result": f"Gemini API Error: {str(e)}"}), 500

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
